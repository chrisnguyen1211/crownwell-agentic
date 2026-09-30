"""Step 5b: summary.xlsx -- summary (event x source), arb_windows, k, and one raw 5m sheet per event.
Style: brown header, white bold text, white background, frozen header row, percentages with 2 decimals."""
from __future__ import annotations

import re

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import timeutil as tu
from .report import PCT, load_event, summary

HEADER_FILL = PatternFill("solid", fgColor="6B4226")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WHITE = PatternFill("solid", fgColor="FFFFFF")


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for c in out.columns:
        if isinstance(out[c].dtype, pd.DatetimeTZDtype):
            out[c] = out[c].dt.tz_convert("UTC").dt.tz_localize(None)
    return out.replace([np.inf, -np.inf], np.nan)


def _write(xw, df: pd.DataFrame, name: str, pct_cols=()):
    name = re.sub(r"[\[\]\*\?/\\:]", "_", name)[:31]
    df = _clean(df)
    df.to_excel(xw, sheet_name=name, index=False)
    ws = xw.sheets[name]
    ws.sheet_view.showGridLines = True
    ws.freeze_panes = "A2"
    for j, col in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=j)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200).tolist()]) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 9), 48)
        is_pct = col in pct_cols
        for i in range(2, len(df) + 2):
            c = ws.cell(row=i, column=j)
            c.fill = WHITE
            if is_pct:
                c.number_format = "0.00%"
            elif isinstance(c.value, float):
                c.number_format = "#,##0.00####"
            elif hasattr(c.value, "year"):
                c.number_format = "yyyy-mm-dd hh:mm"


def build(paths, cfg, event_ids: list[str]) -> str:
    all_d = [d for d in (load_event(paths, e) for e in event_ids) if d]
    met, evs, by_issuer = summary(all_d)
    pct = PCT | {"gap_stock", "gap", "premium_mean_night", "abs_residual_preopen_mean"} | \
        {c for c in met.columns if c.startswith("captured_")}
    with pd.ExcelWriter(paths.xlsx, engine="openpyxl") as xw:
        _write(xw, met, "summary", pct)
        _write(xw, by_issuer, "by_issuer", pct)
        _write(xw, evs, "by_event", pct)
        arb = pd.concat([d["arb"] for d in all_d if not d["arb"].empty], ignore_index=True) \
            if any(not d["arb"].empty for d in all_d) else pd.DataFrame(columns=["event_id"])
        for c in ("start", "end", "peak_at"):
            if c in arb:
                arb[c + "_et"] = arb[c].map(tu.fmt_et)
                arb[c + "_vn"] = arb[c].map(tu.fmt_vn)
                arb[c] = pd.to_datetime(arb[c], utc=True)
        _write(xw, arb, "arb_windows", pct)
        k = pd.concat([d["k"] for d in all_d if not d["k"].empty], ignore_index=True) if all_d else pd.DataFrame()
        _write(xw, k, "k_factors", {"k_minus_1"})
        ll = pd.concat([d["leadlag"] for d in all_d if not d["leadlag"].empty], ignore_index=True) \
            if any(not d["leadlag"].empty for d in all_d) else pd.DataFrame()
        _write(xw, ll, "lead_lag")
        for d in all_d:
            ev = d["ev"]
            m = (d["adj"].index >= pd.Timestamp(ev["win_start"])) & (d["adj"].index < pd.Timestamp(ev["win_end"]))
            lab = d["labels"]
            wide = pd.DataFrame(index=d["adj"].index[m])
            wide["ts_utc"] = wide.index.tz_convert("UTC").tz_localize(None)
            wide["ts_et"] = wide.index.tz_convert(tu.ET).strftime("%Y-%m-%d %H:%M")
            wide["ts_vn"] = wide.index.tz_convert(tu.VN).strftime("%Y-%m-%d %H:%M")
            wide["session"] = [tu.us_session_label(t) for t in wide.index]
            for s in d["adj"].columns:
                wide[f"{lab[s]} P_adj"] = d["adj"].loc[m, s]
            for s in d["volq"].columns:
                if d["volq"][s].notna().any():
                    wide[f"{lab[s]} vol_usd"] = d["volq"].loc[m, s]
            if "dispersion" in d["xs"]:
                wide["median_price"] = d["xs"]["median_price"]
                wide["dispersion"] = d["xs"]["dispersion"]
            pcols = {"dispersion"}
            for s in d["prem"].columns:
                wide[f"{lab[s]} premium"] = d["prem"][s]
                pcols.add(f"{lab[s]} premium")
            _write(xw, wide.reset_index(drop=True), f"raw_{ev['ticker']}_{ev['prev_session'][5:]}", pcols)
    return str(paths.xlsx)
